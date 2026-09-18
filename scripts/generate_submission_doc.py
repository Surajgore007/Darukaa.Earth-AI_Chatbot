"""
Darukaa.Earth Submission Document Generator.

Generates the exact Word document (.docx) required by the Darukaa.Earth Hackathon guidelines:
1. GitHub repository link
2. Live demo URL
3. Brief README.md overview (architecture, database/schema, local setup, CI/CD)
4. Other links, credentials, and notes
5. Private repository access invitations list
"""

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_submission_doc(output_path: str = "Darukaa_Earth_Submission.docx"):
    doc = Document()

    # Document Title
    title = doc.add_heading("Darukaa.Earth — AI Biodiversity Intelligence Chatbot Challenge", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph("Submission Document for Internship / Hackathon Challenge")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(12)
    sub.runs[0].font.italic = True

    doc.add_paragraph()

    # Section 1: Project Links & Submission Overview
    doc.add_heading("1. Project Links & Submission Overview", level=1)
    
    p = doc.add_paragraph()
    p.add_run("GitHub Repository Link: ").bold = True
    p.add_run("https://github.com/Surajgore007/Darukaa.Earth-AI_Chatbot.git\n")
    p.add_run("Repository Access: ").bold = True
    p.add_run("Public (accessible without invitation)\n")
    p.add_run("Live Demo URL: ").bold = True
    p.add_run("https://darukaa-earth-ai-chatbot.vercel.app (FastAPI Interactive Web Interface & API Docs at /docs)\n")
    p.add_run("Candidate Name: ").bold = True
    p.add_run("Suraj Gore\n")
    p.add_run("Project: ").bold = True
    p.add_run("Darukaa.Earth — AI Biodiversity Intelligence Chatbot\n")

    # Section 2: Repository Access (for reviewer reference)
    doc.add_heading("2. Repository Access Instructions", level=1)
    doc.add_paragraph(
        "The repository is configured as PUBLIC at https://github.com/Surajgore007/Darukaa.Earth-AI_Chatbot.git. "
        "Reviewers can clone and run it directly. For reference, access permissions are also configured for:"
    )
    accounts = [
        "ankita.dasgupta@darukaa.com",
        "harsh.kumar@darukaa.com",
        "utkarsh.gauniyal@darukaa.com",
        "guneet.mutreja@darukaa.com"
    ]
    for acc in accounts:
        doc.add_paragraph(acc, style="List Bullet")

    # Section 3: Architecture & System Overview
    doc.add_heading("3. System Architecture & Design", level=1)
    doc.add_paragraph(
        "The system operates as an 'AI Environmental Scientist' adhering strictly to the pipeline:"
    )
    doc.add_paragraph(
        "FastAPI → LangGraph → Variable Extraction (Turn Memory) → Completeness Check "
        "→ Dual-Layer Retrieval (pgvector + Structured SQL) → Gemini Reasoning "
        "→ Grounding Validator (Strict Range Check with Retry) → Programmatic Confidence Scorer → Final Response"
    )

    doc.add_heading("Key Technical Highlights:", level=2)
    highlights = [
        ("Multi-Metric Synthesis: ", "Connects multiple variables (soil condition, rainfall/climate, land use, and pollution) into ONE primary integrated intervention rather than producing an overwhelming shopping list of separate practices."),
        ("Verified Scientific Data: ", "Curated genuine datasets from FAO, IPCC (SRCCL), IPBES, and peer-reviewed journals. Zero fabricated studies, statistics, effect sizes, or URLs."),
        ("Dual-Layer Retrieval: ", "Separates qualitative scientific context (knowledge_chunks via pgvector) from quantitative effect sizes (metric_facts via SQL) so numeric claims can be audited independently."),
        ("Strict Grounding Validator: ", "Requires complete numeric expressions/ranges (e.g. +15–25%) to match retrieved evidence with zero loose-digit fallback. Hallucinations trigger an automated strict retry; repeated failure returns 'Insufficient evidence for a confident recommendation.'"),
        ("Programmatic Confidence: ", "Calculated by backend logic (HIGH, MEDIUM, LOW) requiring a verified (intervention -> metric) pair match. Gemini never self-evaluates confidence."),
        ("Exact-Metric Question Rule: ", "If a user requests an exact percentage for a metric not present in evidence (e.g. biodiversity increase percentage), the system explicitly reports that the exact percentage cannot be reported and does not substitute unrelated numbers."),
        ("Verified Gemini Models: ", "Updated to Google's active supported models: gemini-2.5-flash for reasoning and gemini-embedding-001 (768 dimensions) for vector embeddings.")
    ]
    for bold_prefix, text in highlights:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(bold_prefix).bold = True
        p.add_run(text)

    # Section 4: Database Schema
    doc.add_heading("4. Database Schema (Supabase PostgreSQL + pgvector)", level=1)
    doc.add_paragraph(
        "1. Table 'knowledge_chunks':\n"
        "   - id: UUID (Primary Key)\n"
        "   - content: TEXT\n"
        "   - embedding: VECTOR(768) (pgvector IVFFlat cosine similarity index)\n"
        "   - source: TEXT (e.g. FAO, IPCC)\n"
        "   - source_url: TEXT\n"
        "   - variable_tag: TEXT (soil, land_use, biodiversity, climate, human_impact)\n\n"
        "2. Table 'metric_facts':\n"
        "   - id: UUID (Primary Key)\n"
        "   - intervention: TEXT\n"
        "   - affects_metric: TEXT\n"
        "   - effect_value: TEXT (e.g. '+15–25% over 2–3 years')\n"
        "   - time_horizon: TEXT (short, medium, long)\n"
        "   - source: TEXT\n"
        "   - source_url: TEXT"
    )

    # Section 5: Sample System Outputs
    doc.add_heading("5. Sample System Outputs", level=1)
    doc.add_heading("Standard Multi-Variable Output:", level=2)
    sample_text = (
        "RECOMMENDATION:\n"
        "Establish an agroforestry system with integrated legume cover cropping.\n\n"
        "WHY IT WORKS:\n"
        "Depleted soil organic carbon (0.3%) combined with low rainfall and continuous monoculture wheat limits soil microbial "
        "biomass and moisture retention. Introducing nitrogen-fixing trees and legumes creates canopy shade that buffers surface temperatures "
        "and increases organic matter. The physical presence of deep tree roots and cover crop residue improves water infiltration and "
        "reduces reliance on synthetic nitrogen fertilizer, directly mitigating agricultural runoff into local waterways.\n\n"
        "IMPACTED METRICS:\n"
        "- Soil moisture retention in root zone: +20-40% retention in root zone\n"
        "- Microbial biomass: +20-35% increase in microbial community activity\n\n"
        "TIME HORIZON: medium\n"
        "CONFIDENCE: HIGH\n"
        "SOURCES:\n"
        "- FAO - State of Knowledge of Soil Biodiversity (2020) | https://www.fao.org/documents/card/en/c/cb1928en\n"
        "- IPCC Special Report on Climate Change and Land (SRCCL), Chapter 4 (2019) | https://www.ipcc.ch/srccl/chapter/chapter-4/"
    )
    doc.add_paragraph(sample_text)

    doc.add_heading("Handling Exact Unsupported Metric Queries:", level=2)
    doc.add_paragraph(
        "User Query: 'By exactly what percentage will biodiversity increase from agroforestry?'\n"
        "System Output: 'The retrieved evidence does not provide an exact percentage for biodiversity improvement from agroforestry, "
        "so an exact biodiversity percentage cannot be reported. However, retrieved evidence does substantiate related soil moisture "
        "retention (+20-40%) and carbon benefits.' (Confidence: LOW)"
    )

    # Section 6: Local Setup & Testing
    doc.add_heading("6. Local Setup & Testing Instructions", level=1)
    doc.add_paragraph(
        "1. Python Environment:\n"
        "   - py -3.11 -m venv .venv\n"
        "   - .venv\\Scripts\\activate\n"
        "   - pip install -r requirements.txt\n\n"
        "2. Environment Configuration (.env):\n"
        "   - GEMINI_API_KEY=your_key\n"
        "   - DATABASE_URL=postgresql://postgres:[pass]@db.[project].supabase.co:5432/postgres\n\n"
        "3. Automated Knowledge Ingestion:\n"
        "   - python scripts/ingest_knowledge.py\n\n"
        "4. Start Application:\n"
        "   - python -m uvicorn app.main:app --port 8000\n"
        "   - Open browser at: http://localhost:8000\n\n"
        "5. Automated Test Suite:\n"
        "   - pytest -v (Full suite of 33 unit and integration tests — 100% pass rate)\n"
        "   - python tests/test_queries.py (Runs 9 challenge benchmark scenarios)\n"
        "   - pytest tests/test_submission_scenarios.py -v (Runs 6 submission scenarios A–F)"
    )

    # Section 7: CI/CD Pipeline
    doc.add_heading("7. CI/CD & Automated Quality Checks", level=1)
    doc.add_paragraph(
        "A GitHub Actions CI/CD workflow (.github/workflows/ci.yml) runs on every push and pull request. "
        "It automatically installs dependencies, executes all 33 pytest tests, and runs the 9 benchmark scenarios "
        "to guarantee continuous grounding integrity."
    )

    doc.save(output_path)
    print(f"[SUCCESS] Submission document generated at: {output_path}")


if __name__ == "__main__":
    create_submission_doc()
